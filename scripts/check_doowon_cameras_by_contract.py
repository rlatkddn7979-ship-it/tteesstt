#!/usr/bin/env python3
"""
어떤 회사든 "계약업체명으로 계약번호부터 찾고, 그 계약번호로 전체 물품을 가져오는" 2단계
방식으로 조회하는 스크립트. check_doowon_cameras_shoppingmall.py를 품명 없이(전국) 돌리면
90만 건이 넘는 데이터를 다 훑어야 해서 비효율적인데, 이 스크립트는 그럴 필요가 없다.

동작 방식:
  1단계) getThptyUcntrctPrdctInfoList를 계약업체명(cntrctCorpNm)=회사명으로 조회해서
         (빠르고 좁은 조회) 그 회사가 가진 쇼핑계약번호(shopngCntrctNo) 목록을 알아낸다.
  2단계) 알아낸 계약번호 각각에 대해 getShoppingMallPrdctInfoList를 조회구분(inqryDiv)=2,
         쇼핑계약번호(shopngCntrctNo)로 직접 조회한다. 이러면 그 계약에 딸린 물품을
         (제3자단가계약의 기본 등록 품목이든, 나중에 추가된 "추가선택품목"이든) 전부 가져올
         수 있어서, 날짜 범위 스캔이나 전국 데이터를 다 훑을 필요가 없다.

그래서 회사명만 바꿔가며 어떤 업체든 빠르게(전국 스캔 없이) 전체 등록 물품을 확인할 수 있다.

주의: 1단계에서 계약업체명으로 걸리는 계약번호만 찾기 때문에, "제조사만 이 회사고
계약업체명은 다른 회사"인 경우(check_doowon_cameras_bymaker.py가 다루는 경우)는 이 방식으로
못 찾는다. 그런 경우까지 찾고 싶으면 bymaker나 corp_or_maker 스크립트를 같이 써야 한다.

data.go.kr 활용신청: https://www.data.go.kr/data/15129471/openapi.do
Endpoint: https://apis.data.go.kr/1230000/at/ShoppingMallPrdctInfoService

사용법 (CLI):
    export DATA_GO_KR_SERVICE_KEY="발급받은 서비스키(디코딩된 일반 인증키)"
    python3 scripts/check_doowon_cameras_by_contract.py --corp-name "두원전자통신" --category ""
    python3 scripts/check_doowon_cameras_by_contract.py --corp-name "다른회사이름"

GUI로 실행하려면 scripts/check_doowon_cameras_by_contract_gui.py 를 실행하세요.

서비스키는 절대 코드에 하드코딩하거나 커밋하지 마세요. 환경변수로만 주입합니다.
"""

import argparse
import collections
import csv
import datetime
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import uuid

BASE_URL = "https://apis.data.go.kr/1230000/at/ShoppingMallPrdctInfoService"
STAGE1_OPERATION = "getThptyUcntrctPrdctInfoList"  # 제3자계약 품목정보 조회 (계약번호 알아내기용)
STAGE2_OPERATION = "getShoppingMallPrdctInfoList"  # 나라장터 쇼핑몰 품목 정보 목록 조회 (전체 물품)

if getattr(sys, "frozen", False):
    # PyInstaller 등으로 실행파일(.exe)로 묶인 경우, __file__은 임시 압축해제 폴더를
    # 가리켜서 그 기준으로 저장하면 프로그램 종료 후 파일이 사라진다. 이때는 실행파일
    # 자체가 있는 폴더를 기준으로 삼는다.
    SCRIPT_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def _resolve_output_path(path: str) -> str:
    return path if os.path.isabs(path) else os.path.join(SCRIPT_DIR, path)

MAX_ITEMS = 20000  # 안전장치 (정상적인 회사 하나의 물품 수는 이보다 훨씬 적어야 함)

NAME_FIELD = "prdctClsfcNoNm"  # 품명
SPEC_FIELD = "prdctSpecNm"  # 규격/모델명
CORP_FIELD = "cntrctCorpNm"  # 계약업체명
CONTRACT_FIELD = "shopngCntrctNo"  # 쇼핑계약번호

EXPORT_COLUMNS = [
    ("계약업체명", "cntrctCorpNm"),
    ("제조사", "prdctMakrNm"),
    ("품명", "prdctClsfcNoNm"),
    ("세부품명", "dtilPrdctClsfcNoNm"),
    ("규격/모델명", "prdctSpecNm"),
    ("물품식별번호", "prdctIdntNo"),
    ("계약방법", "cntrctMthdNm"),
    ("쇼핑계약번호", "shopngCntrctNo"),
    ("쇼핑계약순번", "shopngCntrctSno"),
    ("계약일자", "cntrctDate"),
    ("계약시작일", "cntrctBgnDate"),
    ("계약종료일", "cntrctEndDate"),
    ("계약금액", "cntrctPrceAmt"),
    ("단위", "prdctUnit"),
    ("등록일시", "rgstDt"),
]


CURL_STATUS_MARKER = "__HTTP_STATUS__"


def _curl_request(url: str, timeout: int):
    """curl.exe로 요청을 보낸다. 이 PC에서는 파이썬 urllib은 계속 타임아웃 나는데
    curl은 항상 정상 응답을 받아온 게 확인돼서, urllib 대신 curl을 직접 호출한다."""
    try:
        result = subprocess.run(
            [
                "curl", "-s", "-S",
                "--max-time", str(timeout),
                "-A", "curl/8.5.0",
                "-H", "Accept: application/json",
                "-w", f"\n{CURL_STATUS_MARKER}%{{http_code}}",
                url,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=timeout + 10,
            stdin=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired as e:
        raise TimeoutError(f"curl 실행이 {timeout + 10}초를 넘겨 중단됨") from e
    except FileNotFoundError as e:
        raise RuntimeError(
            "curl 명령을 찾을 수 없습니다. Windows 10/11에는 기본 내장되어 있는데, "
            "PATH에서 안 잡히면 PowerShell에서 'curl' 실행이 되는지 먼저 확인해주세요."
        ) from e

    if result.returncode != 0:
        raise TimeoutError(f"curl 오류(exit code {result.returncode}): {result.stderr.strip()}")

    output = result.stdout
    body, _, status_part = output.rpartition(CURL_STATUS_MARKER)
    status_code = int(status_part.strip()) if status_part.strip().isdigit() else None
    return body, status_code


def fetch(
    operation: str,
    service_key: str,
    params: dict,
    timeout: int = 60,
    retries: int = 10,
    log=print,
) -> dict:
    query = dict(params)
    query["serviceKey"] = service_key
    url = f"{BASE_URL}/{operation}?{urllib.parse.urlencode(query)}"
    log(f"요청 URL: {url}")

    RETRYABLE_HTTP_CODES = (502, 503, 504)

    last_error = None
    for attempt in range(1, retries + 1):
        try:
            raw, status_code = _curl_request(url, timeout)
        except TimeoutError as e:
            last_error = e
            log(f"  (타임아웃, {attempt}/{retries}번째 시도 실패: {e})")
        else:
            if status_code is None or status_code < 400:
                break
            last_error = urllib.error.HTTPError(url, status_code, raw[:200], None, None)
            if status_code not in RETRYABLE_HTTP_CODES:
                raise last_error
            log(f"  (HTTP {status_code}, {attempt}/{retries}번째 시도 실패)")

        if attempt < retries:
            log("  대기 없이 바로 재시도합니다...")
    else:
        raise last_error

    try:
        return {"raw": raw, "json": json.loads(raw)}
    except json.JSONDecodeError:
        return {"raw": raw, "json": None}


def extract_items(payload: dict):
    """data.go.kr 표준 응답 포맷(response.body.items)에서 item 리스트를 뽑아본다."""
    if not payload:
        return []
    body = (
        payload.get("response", {}).get("body", {})
        if isinstance(payload.get("response"), dict)
        else None
    )
    if not body:
        return []
    items = body.get("items")
    if isinstance(items, dict):
        item = items.get("item", [])
        return item if isinstance(item, list) else [item]
    if isinstance(items, list):
        return items
    return []


def _date_chunks(begin_date: str, end_date: str, max_days: int = 366):
    """rgstDtBgnDt/rgstDtEndDt는 한 번에 최대 1년까지만 조회 가능하므로,
    [begin_date, end_date](YYYYMMDD)를 max_days 이하 구간들로 쪼갠다."""
    begin = datetime.datetime.strptime(begin_date, "%Y%m%d").date()
    end = datetime.datetime.strptime(end_date, "%Y%m%d").date()
    chunks = []
    cur = begin
    while cur <= end:
        chunk_end = min(cur + datetime.timedelta(days=max_days - 1), end)
        chunks.append((cur.strftime("%Y%m%d"), chunk_end.strftime("%Y%m%d")))
        cur = chunk_end + datetime.timedelta(days=1)
    return chunks


def _paginate(operation, base_params, num_of_rows, service_key, log, stage_label):
    """모든 페이지를 넘겨가며 item을 다 모아서 리스트로 반환한다."""
    items_out = []
    page_no = 1
    total_count = None
    while True:
        params = dict(base_params, pageNo=str(page_no))
        try:
            result = fetch(operation, service_key, params, log=log)
        except urllib.error.HTTPError as e:
            log(f"HTTP 오류: {e.code} {e.reason}")
            break
        except urllib.error.URLError as e:
            log(f"네트워크 오류: {e.reason}")
            break
        except TimeoutError as e:
            log(f"타임아웃 (재시도 모두 실패): {e}")
            break

        payload = result["json"]
        if payload is None:
            log("JSON이 아닌 응답 (앞부분 500자):")
            log(result["raw"][:500])
            break

        header = payload.get("response", {}).get("header", {})
        result_code = header.get("resultCode")
        log(f"[{stage_label}][page {page_no}] resultCode={result_code} resultMsg={header.get('resultMsg')}")
        if result_code not in ("00", "0", None):
            break

        body = payload.get("response", {}).get("body", {})
        if total_count is None:
            total_count = body.get("totalCount")
        items = extract_items(payload)
        log(f"[{stage_label}][page {page_no}] 조회된 item 수: {len(items)} (totalCount={total_count})")
        items_out.extend(items)

        if not items:
            break
        if total_count is not None and len(items_out) >= int(total_count):
            break
        if len(items) < num_of_rows:
            break
        if len(items_out) >= MAX_ITEMS:
            log(f"경고: {MAX_ITEMS}건 이상 조회되어 중단합니다.")
            break
        page_no += 1
        time.sleep(1)

    return items_out


def run_query(
    corp_name: str,
    category: str,
    spec: str,
    begin_date: str,
    end_date: str,
    service_key: str,
    num_of_rows: int = 999,
    log=print,
) -> dict:
    """1단계: 계약업체명으로 쇼핑계약번호를 알아내고, 2단계: 그 계약번호로 전체 물품을
    가져온다 (전국 스캔 없이). category(OR)/spec(AND)로 결과를 좁힌다.

    반환값: {"ok": bool, "error": str|None, "all_items": [...],
             "camera_items": [...], "distinct_names": [...], "name_field": str|None}
    """
    date_chunks = _date_chunks(begin_date, end_date)
    if len(date_chunks) > 1:
        log(f"등록일시 범위가 1년을 넘어서 {len(date_chunks)}개 구간으로 나눠 1단계를 조회합니다.")

    log(f"\n### 1단계: 계약업체명(cntrctCorpNm)='{corp_name}' 으로 쇼핑계약번호 찾기 ###")
    stage1_items = []
    for chunk_idx, (chunk_begin, chunk_end) in enumerate(date_chunks, start=1):
        log(f"\n--- 구간 {chunk_idx}/{len(date_chunks)}: {chunk_begin}~{chunk_end} ---")
        base_params = {
            "numOfRows": str(num_of_rows),
            "type": "json",
            "cntrctCorpNm": corp_name,
            "rgstDtBgnDt": f"{chunk_begin}0000",
            "rgstDtEndDt": f"{chunk_end}2359",
        }
        stage1_items.extend(
            _paginate(STAGE1_OPERATION, base_params, num_of_rows, service_key, log, "1단계")
        )
        if chunk_idx < len(date_chunks):
            time.sleep(1)

    contract_numbers = sorted({
        str(item.get(CONTRACT_FIELD, "")).strip()
        for item in stage1_items
        if str(item.get(CONTRACT_FIELD, "")).strip()
    })

    if not contract_numbers:
        log(
            f"\n'{corp_name}'의 쇼핑계약번호를 1단계에서 찾지 못했습니다. "
            "계약업체명 표기가 다르거나(예: '(주)두원전자통신'), 조회 기간 밖의 계약일 수 있습니다. "
            "check_doowon_cameras_bymaker.py나 corp_or_maker.py로도 확인해보세요."
        )
        return {"ok": False, "error": "no_contract_number", "all_items": [],
                "camera_items": [], "distinct_names": [], "name_field": None}

    log(f"\n찾은 쇼핑계약번호 ({len(contract_numbers)}개): {', '.join(contract_numbers)}")

    log(f"\n### 2단계: 각 쇼핑계약번호로 전체 물품 가져오기 (getShoppingMallPrdctInfoList) ###")
    all_items = []
    seen_keys = set()

    def add_items(items):
        new_count = 0
        for item in items:
            key = json.dumps(item, sort_keys=True, ensure_ascii=False)
            if key not in seen_keys:
                seen_keys.add(key)
                all_items.append(item)
                new_count += 1
        return new_count

    add_items(stage1_items)

    for contract_no in contract_numbers:
        log(f"\n--- 쇼핑계약번호: {contract_no} ---")
        base_params = {
            "numOfRows": str(num_of_rows),
            "type": "json",
            "inqryDiv": "2",
            "shopngCntrctNo": contract_no,
        }
        items = _paginate(STAGE2_OPERATION, base_params, num_of_rows, service_key, log, "2단계")
        new_count = add_items(items)
        log(f"  이 계약번호에서 새로 추가된 물품 수: {new_count} (이 계약 전체 {len(items)}건)")

    log(f"\n1+2단계 합산(중복 제거) 전체 물품 수: {len(all_items)}")

    name_field = NAME_FIELD
    category_counts = collections.Counter(str(item.get(name_field, "")) for item in all_items)
    all_categories = sorted(category_counts)

    category_terms = [t.strip() for t in category.split(",") if t.strip()]
    spec_terms = [t.strip() for t in spec.split(",") if t.strip()]

    def matches(item):
        if category_terms:
            category_text = str(item.get(NAME_FIELD, ""))
            if not any(term in category_text for term in category_terms):
                return False
        if spec_terms:
            spec_text = str(item.get(SPEC_FIELD, ""))
            if not all(term in spec_text for term in spec_terms):
                return False
        return True

    camera_items = [item for item in all_items if matches(item)]
    distinct_names = sorted({str(item.get(name_field, "")) for item in camera_items})

    filter_desc = f"품명(OR)={category_terms or '(없음)'}, 규격(AND)={spec_terms or '(없음)'}"

    log(f"\n'{corp_name}' 전체 등록 물품 수: {len(all_items)}")
    log(f"품명({name_field}) 종류 ({len(all_categories)}개):")
    for cat_name in all_categories:
        log(f"  - {cat_name} ({category_counts[cat_name]}건)")

    log(f"\n조건({filter_desc}) 일치 물품 수: {len(camera_items)}")
    if not camera_items:
        log(f"{filter_desc} 조건과 일치하는 물품이 없습니다. 규격 조건을 줄여보세요.")
    log(f"일치하는 물품 종류(고유 {name_field} 개수): {len(distinct_names)}")
    for name in distinct_names:
        log(f"  - {name}")

    return {
        "ok": True,
        "error": None,
        "all_items": all_items,
        "camera_items": camera_items,
        "distinct_names": distinct_names,
        "name_field": name_field,
    }


NUMERIC_EXPORT_FIELDS = {"cntrctPrceAmt"}
ACCOUNTING_NUMBER_FORMAT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'


INVALID_FILENAME_CHARS = '\\/:*?"<>|'


def _sanitize_filename_part(text: str) -> str:
    return "".join(c for c in text if c not in INVALID_FILENAME_CHARS).strip()


def default_output_filename(corp_name: str, category: str, spec: str, ext: str = ".xlsx") -> str:
    """'업체명_품명_규격_실행시각' 형태의 파일명을 만든다 (품명/규격이 비어있으면 생략)."""
    parts = [_sanitize_filename_part(corp_name) or "업체"]
    if category.strip():
        parts.append(_sanitize_filename_part(category.replace(",", "-")))
    if spec.strip():
        parts.append(_sanitize_filename_part(spec.replace(",", "-")))
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    parts.append(timestamp)
    return "_".join(p for p in parts if p) + ext


def _timestamped_path(path):
    root, ext = os.path.splitext(path)
    suffix = uuid.uuid4().hex[:6]
    return f"{root}_retry{suffix}{ext}"


def _write_with_fallback(write_fn, path, log):
    try:
        write_fn(path)
        return path
    except PermissionError:
        fallback = _timestamped_path(path)
        log(
            f"'{path}' 파일을 저장할 수 없어(다른 프로그램에서 사용 중이거나 잠겨 있을 수 있음) "
            f"'{fallback}'(으)로 대신 저장합니다."
        )
        try:
            write_fn(fallback)
        except PermissionError as e:
            raise PermissionError(
                f"'{fallback}'로도 저장할 수 없습니다. 해당 폴더에 쓰기 권한이 없거나 "
                "(예: OneDrive 동기화 중, 보안 프로그램 차단 등) 폴더 자체의 문제일 수 있습니다. "
                "다른 폴더(예: C:\\Temp)로 --output 경로를 바꿔서 시도해보세요."
            ) from e
        return fallback


def write_output(camera_items, corp_name, category, spec, begin_date, end_date, output_path, log=print):
    """물품 목록을 엑셀(.xlsx)로 저장한다. openpyxl이 없으면 CSV로 대신 저장한다."""
    output_path = _resolve_output_path(output_path)
    log(f"저장 경로: {output_path}")

    try:
        import openpyxl
    except ImportError:
        openpyxl = None

    def cell_value(item, field):
        value = item.get(field, "")
        if field in NUMERIC_EXPORT_FIELDS and value not in ("", None):
            try:
                return int(value)
            except (TypeError, ValueError):
                pass
        return value

    rows = [[cell_value(item, field) for _, field in EXPORT_COLUMNS] for item in camera_items]
    headers = [header for header, _ in EXPORT_COLUMNS]

    def save_csv(p):
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(rows)

    def csv_fallback_path():
        return os.path.splitext(output_path)[0] + ".csv"

    if openpyxl is not None:
        path = output_path
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "물품목록"
        ws.append(headers)
        for row in rows:
            ws.append(row)

        numeric_cols = [
            col_idx for col_idx, (_, field) in enumerate(EXPORT_COLUMNS, start=1)
            if field in NUMERIC_EXPORT_FIELDS
        ]
        for col_idx in numeric_cols:
            for row_idx in range(2, len(rows) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = ACCOUNTING_NUMBER_FORMAT

        for col_idx, header in enumerate(headers, start=1):
            width = max(len(header), *(len(str(r[col_idx - 1])) for r in rows)) if rows else len(header)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(width + 2, 60)

        summary = wb.create_sheet("요약")
        summary.append(["업체명", corp_name])
        summary.append(["품명", category])
        summary.append(["규격", spec])
        summary.append(["조회 기간(1단계 계약번호 탐색용)", f"{begin_date} ~ {end_date}"])
        summary.append(["매칭 물품 수", len(camera_items)])

        try:
            saved_path = _write_with_fallback(wb.save, path, log)
        except PermissionError:
            csv_path = csv_fallback_path()
            log(
                f"\n엑셀(.xlsx) 저장이 계속 막혀서 CSV로 대신 저장합니다: {csv_path} "
                "(엑셀 파일 자체가 보안 프로그램에 의해 차단되고 있을 수 있습니다)"
            )
            saved_path = _write_with_fallback(save_csv, csv_path, log)
        log(f"\n결과를 저장했습니다: {saved_path}")
        return saved_path
    else:
        path = csv_fallback_path()
        saved_path = _write_with_fallback(save_csv, path, log)
        log(f"\nopenpyxl이 설치되어 있지 않아 CSV로 저장했습니다: {saved_path}")
        log("엑셀(.xlsx)로 저장하려면 'pip install openpyxl' 실행 후 다시 실행하세요.")
        return saved_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--corp-name", default="두원전자통신", help="조회할 업체명 (계약업체명 기준)")
    parser.add_argument(
        "--category",
        default="",
        help="품명(대분류) 필터. 쉼표로 여러 개를 넣으면 그 중 하나라도(OR) 일치하면 매칭. "
             "비워두면 이 회사의 계약번호에 딸린 물품 전체를 다 보여줌",
    )
    parser.add_argument(
        "--spec",
        default="",
        help="규격/모델 필터. 쉼표로 여러 조건을 넣으면 전부(AND) 포함해야 매칭. 비우면 필터링 안 함",
    )
    parser.add_argument("--num-of-rows", type=int, default=999, help="페이지당 조회 건수")
    parser.add_argument(
        "--begin-date",
        default=(datetime.date.today() - datetime.timedelta(days=365)).strftime("%Y%m%d"),
        help="1단계(계약번호 탐색)에 쓸 조회 시작일자(YYYYMMDD). 기본값: 오늘로부터 1년 전",
    )
    parser.add_argument(
        "--end-date",
        default=datetime.date.today().strftime("%Y%m%d"),
        help="1단계(계약번호 탐색)에 쓸 조회 종료일자(YYYYMMDD)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="결과를 저장할 엑셀 파일 경로. 생략하면 '업체명_품명_규격_실행시각.xlsx'로 자동 생성",
    )
    args = parser.parse_args()

    service_key = os.environ.get("DATA_GO_KR_SERVICE_KEY")
    if not service_key:
        print("환경변수 DATA_GO_KR_SERVICE_KEY 가 설정되어 있지 않습니다.", file=sys.stderr)
        print('예: export DATA_GO_KR_SERVICE_KEY="발급받은 서비스키"', file=sys.stderr)
        sys.exit(1)

    result = run_query(
        args.corp_name,
        args.category,
        args.spec,
        args.begin_date,
        args.end_date,
        service_key,
        num_of_rows=args.num_of_rows,
    )

    if not result["ok"]:
        sys.exit(2)

    output_path = args.output or default_output_filename(args.corp_name, args.category, args.spec)
    write_output(
        result["camera_items"],
        args.corp_name,
        args.category,
        args.spec,
        args.begin_date,
        args.end_date,
        output_path,
    )


if __name__ == "__main__":
    main()
