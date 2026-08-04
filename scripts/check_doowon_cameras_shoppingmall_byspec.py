#!/usr/bin/env python3
"""
check_doowon_cameras_shoppingmall.py와 동일하게 getShoppingMallPrdctInfoList로 조회하되,
마지막 "종류 수" 집계를 품명(prdctClsfcNoNm)이 아니라 규격/모델명(prdctSpecNm) 기준으로 한다.

품명만으로 세면 돔카메라/블렛카메라가 전부 "보안용카메라"라는 같은 품명으로 묶여서 구분이
안 된다. prdctSpecNm에는 보통 "제조사, 모델명, 화소, 형태(돔형/블렛형 등)"가 자유 서술형으로
같이 적혀 있어서, 이 필드를 기준으로 종류를 세면 돔형/블렛형처럼 세부 형태가 다른 물품이
자동으로 서로 다른 종류로 잡힌다. (참고용으로 품명별 개수도 같이 보여준다.)

data.go.kr 활용신청: https://www.data.go.kr/data/15129471/openapi.do
Endpoint: https://apis.data.go.kr/1230000/at/ShoppingMallPrdctInfoService/getShoppingMallPrdctInfoList

사용법 (CLI):
    export DATA_GO_KR_SERVICE_KEY="발급받은 서비스키(디코딩된 일반 인증키)"
    python3 scripts/check_doowon_cameras_shoppingmall_byspec.py
    python3 scripts/check_doowon_cameras_shoppingmall_byspec.py --category "" --begin-date 20200101

GUI로 실행하려면 scripts/check_doowon_cameras_shoppingmall_byspec_gui.py 를 실행하세요.

서비스키는 절대 코드에 하드코딩하거나 커밋하지 마세요. 환경변수로만 주입합니다.
"""

import argparse
import collections
import csv
import datetime
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import uuid

BASE_URL = "https://apis.data.go.kr/1230000/at/ShoppingMallPrdctInfoService"
OPERATION = "getShoppingMallPrdctInfoList"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def _resolve_output_path(path: str) -> str:
    return path if os.path.isabs(path) else os.path.join(SCRIPT_DIR, path)

# 무한루프 방지용 안전장치일 뿐, 실제 사용 상한이 아니다. 품명 없이(전국) 조회하면
# 정상적으로 수십만 건까지 나올 수 있으므로(예: 94만 건), 너무 낮게 잡으면 실제 데이터에
# 도달하기도 전에 조회가 중단되어 원하는 물품을 놓치게 된다.
MAX_ITEMS = 2_000_000

NAME_FIELD = "prdctClsfcNoNm"  # 품명 (참고용)
DTIL_NAME_FIELD = "dtilPrdctClsfcNoNm"  # 세부품명
SPEC_FIELD = "prdctSpecNm"  # 물품규격명 (자유 서술형 설명) - 이 스크립트의 "종류" 집계 기준
CORP_FIELD = "cntrctCorpNm"  # 계약업체명
MAKER_FIELD = "prdctMakrNm"  # 물품제조사명

EXPORT_COLUMNS = [
    ("계약업체명", "cntrctCorpNm"),
    ("제조사", "prdctMakrNm"),
    ("품명", "prdctClsfcNoNm"),
    ("세부품명", "dtilPrdctClsfcNoNm"),
    ("물품규격명", "prdctSpecNm"),
    ("물품식별번호", "prdctIdntNo"),
    ("계약방법", "cntrctMthdNm"),
    ("쇼핑계약번호", "shopngCntrctNo"),
    ("쇼핑계약순번", "shopngCntrctSno"),
    ("계약일자", "cntrctDate"),
    ("계약시작일", "cntrctBgnDate"),
    ("계약종료일", "cntrctEndDate"),
    ("계약가격금액", "cntrctPrceAmt"),
    ("단위", "prdctUnit"),
    ("등록일시", "rgstDt"),
    ("계약업체 사업자등록번호", "cntrctCorpBizno"),
]


CURL_STATUS_MARKER = "__HTTP_STATUS__"


def _curl_request(url: str, timeout: int):
    """curl.exe로 요청을 보낸다. 이 PC에서는 파이썬 urllib은 계속 타임아웃 나는데
    curl은 항상 정상 응답을 받아온 게 확인돼서, urllib 대신 curl을 직접 호출한다."""
    try:
        result = subprocess.run(
            [
                "curl", "-s", "-S",
                "--max-time", str(timeout),
                "-A", "curl/8.5.0",
                "-H", "Accept: application/json",
                "-w", f"\n{CURL_STATUS_MARKER}%{{http_code}}",
                url,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=timeout + 10,
            stdin=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired as e:
        raise TimeoutError(f"curl 실행이 {timeout + 10}초를 넘겨 중단됨") from e
    except FileNotFoundError as e:
        raise RuntimeError(
            "curl 명령을 찾을 수 없습니다. Windows 10/11에는 기본 내장되어 있는데, "
            "PATH에서 안 잡히면 PowerShell에서 'curl' 실행이 되는지 먼저 확인해주세요."
        ) from e

    if result.returncode != 0:
        raise TimeoutError(f"curl 오류(exit code {result.returncode}): {result.stderr.strip()}")

    output = result.stdout
    body, _, status_part = output.rpartition(CURL_STATUS_MARKER)
    status_code = int(status_part.strip()) if status_part.strip().isdigit() else None
    return body, status_code


def fetch(
    operation: str,
    service_key: str,
    params: dict,
    timeout: int = 60,
    retries: int = 10,
    log=print,
) -> dict:
    query = dict(params)
    query["serviceKey"] = service_key
    url = f"{BASE_URL}/{operation}?{urllib.parse.urlencode(query)}"
    log(f"요청 URL: {url}")

    RETRYABLE_HTTP_CODES = (502, 503, 504)

    last_error = None
    for attempt in range(1, retries + 1):
        try:
            raw, status_code = _curl_request(url, timeout)
        except TimeoutError as e:
            last_error = e
            log(f"  (타임아웃, {attempt}/{retries}번째 시도 실패: {e})")
        else:
            if status_code is None or status_code < 400:
                break
            last_error = urllib.error.HTTPError(url, status_code, raw[:200], None, None)
            if status_code not in RETRYABLE_HTTP_CODES:
                raise last_error
            log(f"  (HTTP {status_code}, {attempt}/{retries}번째 시도 실패)")

        if attempt < retries:
            log("  대기 없이 바로 재시도합니다...")
    else:
        raise last_error

    try:
        return {"raw": raw, "json": json.loads(raw)}
    except json.JSONDecodeError:
        return {"raw": raw, "json": None}


def extract_items(payload: dict):
    """data.go.kr 표준 응답 포맷(response.body.items)에서 item 리스트를 뽑아본다."""
    if not payload:
        return []
    body = (
        payload.get("response", {}).get("body", {})
        if isinstance(payload.get("response"), dict)
        else None
    )
    if not body:
        return []
    items = body.get("items")
    if isinstance(items, dict):
        item = items.get("item", [])
        return item if isinstance(item, list) else [item]
    if isinstance(items, list):
        return items
    return []


def _date_chunks(begin_date: str, end_date: str, max_days: int = 366):
    """inqryBgnDate/inqryEndDate는 한 번에 최대 12개월까지만 조회 가능하므로,
    [begin_date, end_date](YYYYMMDD)를 max_days 이하 구간들로 쪼갠다."""
    begin = datetime.datetime.strptime(begin_date, "%Y%m%d").date()
    end = datetime.datetime.strptime(end_date, "%Y%m%d").date()
    chunks = []
    cur = begin
    while cur <= end:
        chunk_end = min(cur + datetime.timedelta(days=max_days - 1), end)
        chunks.append((cur.strftime("%Y%m%d"), chunk_end.strftime("%Y%m%d")))
        cur = chunk_end + datetime.timedelta(days=1)
    return chunks


def run_query(
    corp_name: str,
    category: str,
    spec: str,
    begin_date: str,
    end_date: str,
    service_key: str,
    num_of_rows: int = 999,
    log=print,
) -> dict:
    """getShoppingMallPrdctInfoList로 조회하고, 계약업체명 또는 제조사에 corp_name이
    들어있는 물품을 모두(OR) 찾는다. category(품명)를 지정하면 서버 조회 범위를 그
    품명으로 좁히고, 비워두면 전국 데이터를 다 받아온다(느림).

    check_doowon_cameras_shoppingmall.py와 다른 점: "종류 수"(distinct_names)를
    품명(prdctClsfcNoNm)이 아니라 규격/모델명(prdctSpecNm) 기준으로 센다. 그래야
    "보안용카메라"라는 같은 품명 안에서도 돔형/블렛형처럼 규격이 다른 물품을
    서로 다른 종류로 구분할 수 있다. 품명별 개수는 참고용으로 같이 보여준다.

    반환값: {"ok": bool, "error": str|None, "all_items": [...],
             "camera_items": [...], "distinct_names": [...], "name_field": str|None}
    """
    date_chunks = _date_chunks(begin_date, end_date)
    if len(date_chunks) > 1:
        log(f"등록일자 범위가 12개월을 넘어서 {len(date_chunks)}개 구간으로 나눠 조회합니다.")

    category_terms = [t.strip() for t in category.split(",") if t.strip()]
    server_terms = category_terms or [""]
    if not category_terms:
        log(
            "\n경고: 품명(--category)을 지정하지 않아 전국 모든 회사의 데이터를 조회합니다. "
            "시간이 오래 걸리고 서버 부담이 클 수 있으니 가능하면 품명을 지정해주세요."
        )

    all_items = []
    seen_keys = set()
    any_success = False

    for term in server_terms:
        term_desc = term or "(품명 필터 없음)"
        for chunk_idx, (chunk_begin, chunk_end) in enumerate(date_chunks, start=1):
            log(
                f"\n=== {OPERATION} 품명='{term_desc}' "
                f"(구간 {chunk_idx}/{len(date_chunks)}: {chunk_begin}~{chunk_end}) ==="
            )
            base_params = {
                "numOfRows": str(num_of_rows),
                "type": "json",
                "inqryDiv": "1",
                "inqryBgnDate": chunk_begin,
                "inqryEndDate": chunk_end,
            }
            if term:
                base_params["prdctClsfcNoNm"] = term
            page_no = 1
            total_count = None
            chunk_item_count = 0
            while True:
                params = dict(base_params, pageNo=str(page_no))
                try:
                    result = fetch(OPERATION, service_key, params, log=log)
                except urllib.error.HTTPError as e:
                    log(f"HTTP 오류: {e.code} {e.reason}")
                    break
                except urllib.error.URLError as e:
                    log(f"네트워크 오류: {e.reason}")
                    break
                except TimeoutError as e:
                    log(f"타임아웃 (재시도 모두 실패): {e}")
                    break

                payload = result["json"]
                if payload is None:
                    log("JSON이 아닌 응답 (앞부분 500자):")
                    log(result["raw"][:500])
                    break

                header = payload.get("response", {}).get("header", {})
                result_code = header.get("resultCode")
                result_msg = header.get("resultMsg")
                log(f"[page {page_no}] resultCode={result_code} resultMsg={result_msg}")

                if result_code not in ("00", "0", None):
                    break

                any_success = True
                body = payload.get("response", {}).get("body", {})
                if total_count is None:
                    total_count = body.get("totalCount")
                items = extract_items(payload)
                log(f"[page {page_no}] 조회된 item 수: {len(items)} (totalCount={total_count})")
                chunk_item_count += len(items)

                new_count = 0
                for item in items:
                    key = json.dumps(item, sort_keys=True, ensure_ascii=False)
                    if key not in seen_keys:
                        seen_keys.add(key)
                        all_items.append(item)
                        new_count += 1
                if new_count < len(items):
                    log(f"  (중복 {len(items) - new_count}건 제외)")

                if not items:
                    break
                if total_count is not None and chunk_item_count >= int(total_count):
                    break
                if len(items) < num_of_rows:
                    break
                if len(all_items) >= MAX_ITEMS:
                    log(f"경고: {MAX_ITEMS}건 이상 조회되어 중단합니다. --category로 범위를 좁혀주세요.")
                    break
                page_no += 1
                time.sleep(1)

            if chunk_idx < len(date_chunks):
                time.sleep(1)

    if not any_success:
        log(
            "\n정상 응답(item 포함)을 받지 못했습니다. "
            "resultCode/resultMsg를 참고해서 파라미터를 조정해주세요."
        )
        return {"ok": False, "error": "no_working_operation", "all_items": [],
                "camera_items": [], "distinct_names": [], "name_field": None, "raw_items": all_items}

    log(f"\n품명 필터로 조회된 전체 물품 수(업체명 필터링 전): {len(all_items)}")

    def corp_or_maker_match(item):
        corp_text = str(item.get(CORP_FIELD, ""))
        maker_text = str(item.get(MAKER_FIELD, ""))
        return corp_name in corp_text or corp_name in maker_text

    before = len(all_items)
    matched_items = [item for item in all_items if corp_or_maker_match(item)]
    log(
        f"\n계약업체명 또는 제조사에 '{corp_name}' 포함 여부로 필터링: "
        f"{before}건 -> {len(matched_items)}건"
    )

    if not matched_items:
        log(f"'{corp_name}'이(가) 계약업체명 또는 제조사로 들어간 물품이 없습니다.")
        log("표기가 다를 수 있습니다(예: '(주)두원전자통신' 등). 다른 표기로 다시 시도해보세요.")
        return {"ok": False, "error": "no_matching_corp_or_maker", "all_items": [],
                "camera_items": [], "distinct_names": [], "name_field": None, "raw_items": all_items}

    # 참고용: 품명(대분류) 기준 개수
    name_field_counts = collections.Counter(str(item.get(NAME_FIELD, "")) for item in matched_items)
    all_names = sorted(name_field_counts)

    # 이 스크립트의 핵심: "종류 수"는 규격/모델명(prdctSpecNm) 기준으로 센다.
    name_field = SPEC_FIELD
    spec_counts = collections.Counter(str(item.get(SPEC_FIELD, "")) for item in matched_items)

    spec_terms = [t.strip() for t in spec.split(",") if t.strip()]

    def matches(item):
        if category_terms:
            category_text = str(item.get(NAME_FIELD, ""))
            if not any(term in category_text for term in category_terms):
                return False
        if spec_terms:
            spec_text = str(item.get(SPEC_FIELD, ""))
            if not all(term in spec_text for term in spec_terms):
                return False
        return True

    camera_items = [item for item in matched_items if matches(item)]
    distinct_names = sorted({str(item.get(name_field, "")) for item in camera_items})

    filter_desc = (
        f"계약업체명 또는 제조사={corp_name}, "
        f"품명(OR)={category_terms or '(없음)'}, 규격(AND)={spec_terms or '(없음)'}"
    )

    log(f"\n'{corp_name}'(계약업체명 또는 제조사) 전체 물품 수: {len(matched_items)}")
    log(f"[참고용] 품명({NAME_FIELD}) 기준 개수 ({len(all_names)}개):")
    for cat_name in all_names:
        log(f"  - {cat_name} ({name_field_counts[cat_name]}건)")

    log(f"\n조건({filter_desc}) 일치 물품 수: {len(camera_items)}")
    if not camera_items:
        log(f"{filter_desc} 조건과 일치하는 물품이 없습니다. 규격 조건을 줄여보세요.")
    log(f"일치하는 물품 종류(고유 {name_field}=규격/모델명 개수): {len(distinct_names)}")
    for name in distinct_names:
        log(f"  - {name} ({spec_counts.get(name, 0)}건)")

    return {
        "ok": True,
        "error": None,
        "all_items": matched_items,
        "camera_items": camera_items,
        "distinct_names": distinct_names,
        "name_field": name_field,
        "raw_items": all_items,
    }


NUMERIC_EXPORT_FIELDS = {"cntrctPrceAmt"}
ACCOUNTING_NUMBER_FORMAT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'


INVALID_FILENAME_CHARS = '\\/:*?"<>|'


def _sanitize_filename_part(text: str) -> str:
    return "".join(c for c in text if c not in INVALID_FILENAME_CHARS).strip()


def default_output_filename(corp_name: str, category: str, spec: str, ext: str = ".xlsx") -> str:
    """'업체명_품명_규격_실행시각' 형태의 파일명을 만든다 (품명/규격이 비어있으면 생략)."""
    parts = [_sanitize_filename_part(corp_name) or "업체"]
    if category.strip():
        parts.append(_sanitize_filename_part(category.replace(",", "-")))
    if spec.strip():
        parts.append(_sanitize_filename_part(spec.replace(",", "-")))
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    parts.append(timestamp)
    return "_".join(p for p in parts if p) + ext


def _timestamped_path(path):
    root, ext = os.path.splitext(path)
    suffix = uuid.uuid4().hex[:6]
    return f"{root}_retry{suffix}{ext}"


def _write_with_fallback(write_fn, path, log):
    try:
        write_fn(path)
        return path
    except PermissionError:
        fallback = _timestamped_path(path)
        log(
            f"'{path}' 파일을 저장할 수 없어(다른 프로그램에서 사용 중이거나 잠겨 있을 수 있음) "
            f"'{fallback}'(으)로 대신 저장합니다."
        )
        try:
            write_fn(fallback)
        except PermissionError as e:
            raise PermissionError(
                f"'{fallback}'로도 저장할 수 없습니다. 해당 폴더에 쓰기 권한이 없거나 "
                "(예: OneDrive 동기화 중, 보안 프로그램 차단 등) 폴더 자체의 문제일 수 있습니다. "
                "다른 폴더(예: C:\\Temp)로 --output 경로를 바꿔서 시도해보세요."
            ) from e
        return fallback


def write_output(camera_items, corp_name, category, spec, begin_date, end_date, output_path, log=print):
    """물품 목록을 엑셀(.xlsx)로 저장한다. openpyxl이 없으면 CSV로 대신 저장한다."""
    output_path = _resolve_output_path(output_path)
    log(f"저장 경로: {output_path}")

    try:
        import openpyxl
    except ImportError:
        openpyxl = None

    def cell_value(item, field):
        value = item.get(field, "")
        if field in NUMERIC_EXPORT_FIELDS and value not in ("", None):
            try:
                return int(value)
            except (TypeError, ValueError):
                pass
        return value

    rows = [[cell_value(item, field) for _, field in EXPORT_COLUMNS] for item in camera_items]
    headers = [header for header, _ in EXPORT_COLUMNS]

    def save_csv(p):
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(rows)

    def csv_fallback_path():
        return os.path.splitext(output_path)[0] + ".csv"

    if openpyxl is not None:
        path = output_path
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "물품목록"
        ws.append(headers)
        for row in rows:
            ws.append(row)

        numeric_cols = [
            col_idx for col_idx, (_, field) in enumerate(EXPORT_COLUMNS, start=1)
            if field in NUMERIC_EXPORT_FIELDS
        ]
        for col_idx in numeric_cols:
            for row_idx in range(2, len(rows) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = ACCOUNTING_NUMBER_FORMAT

        for col_idx, header in enumerate(headers, start=1):
            width = max(len(header), *(len(str(r[col_idx - 1])) for r in rows)) if rows else len(header)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(width + 2, 60)

        summary = wb.create_sheet("요약")
        summary.append(["업체명(계약업체명 또는 제조사)", corp_name])
        summary.append(["품명", category])
        summary.append(["규격", spec])
        summary.append(["조회 기간", f"{begin_date} ~ {end_date}"])
        summary.append(["매칭 물품 수", len(camera_items)])

        try:
            saved_path = _write_with_fallback(wb.save, path, log)
        except PermissionError:
            csv_path = csv_fallback_path()
            log(
                f"\n엑셀(.xlsx) 저장이 계속 막혀서 CSV로 대신 저장합니다: {csv_path} "
                "(엑셀 파일 자체가 보안 프로그램에 의해 차단되고 있을 수 있습니다)"
            )
            saved_path = _write_with_fallback(save_csv, csv_path, log)
        log(f"\n결과를 저장했습니다: {saved_path}")
        return saved_path
    else:
        path = csv_fallback_path()
        saved_path = _write_with_fallback(save_csv, path, log)
        log(f"\nopenpyxl이 설치되어 있지 않아 CSV로 저장했습니다: {saved_path}")
        log("엑셀(.xlsx)로 저장하려면 'pip install openpyxl' 실행 후 다시 실행하세요.")
        return saved_path


def write_raw_output(raw_items, output_path, log=print):
    """업체명/제조사 필터링 전, 조회된 원본 물품 전체를 CSV로 저장한다.

    수십만 건 규모가 될 수 있어서 openpyxl(.xlsx)은 메모리/속도 부담이 크므로
    항상 CSV로 저장한다 (엑셀에서도 그대로 열어서 필터/정렬할 수 있음).
    """
    output_path = _resolve_output_path(output_path)
    if not output_path.lower().endswith(".csv"):
        output_path = os.path.splitext(output_path)[0] + ".csv"
    log(f"\n원본 데이터({len(raw_items)}건) 저장 경로: {output_path}")

    headers = [header for header, _ in EXPORT_COLUMNS]

    def save_csv(p):
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for item in raw_items:
                writer.writerow([item.get(field, "") for _, field in EXPORT_COLUMNS])

    saved_path = _write_with_fallback(save_csv, output_path, log)
    log(f"원본 데이터를 저장했습니다: {saved_path}")
    return saved_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--corp-name", default="두원전자통신", help="계약업체명 또는 제조사로 찾을 이름")
    parser.add_argument(
        "--category",
        default="보안용카메라",
        help="품명(대분류) 필터. 쉼표로 여러 개를 넣으면 그 중 하나라도(OR) 일치하면 매칭. "
             "비우면 전국 데이터를 다 받아옴(느림)",
    )
    parser.add_argument(
        "--spec",
        default="",
        help="규격/모델 필터. 쉼표로 여러 조건을 넣으면 전부(AND) 포함해야 매칭. 비우면 필터링 안 함",
    )
    parser.add_argument("--num-of-rows", type=int, default=999, help="페이지당 조회 건수")
    parser.add_argument(
        "--begin-date",
        default=(datetime.date.today() - datetime.timedelta(days=365)).strftime("%Y%m%d"),
        help="조회 시작일자(YYYYMMDD). 기본값: 오늘로부터 1년 전",
    )
    parser.add_argument(
        "--end-date",
        default=datetime.date.today().strftime("%Y%m%d"),
        help="조회 종료일자(YYYYMMDD)",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="결과를 저장할 엑셀 파일 경로. 생략하면 '업체명_품명_규격_실행시각.xlsx'로 자동 생성",
    )
    parser.add_argument(
        "--save-raw",
        action="store_true",
        help="업체명/제조사로 걸러내기 전, 조회된 원본 물품 전체도 별도 CSV로 저장 "
             "(품명 없이 조회하면 수십만 건이 될 수 있어 항상 CSV로 저장됨)",
    )
    args = parser.parse_args()

    service_key = os.environ.get("DATA_GO_KR_SERVICE_KEY")
    if not service_key:
        print("환경변수 DATA_GO_KR_SERVICE_KEY 가 설정되어 있지 않습니다.", file=sys.stderr)
        print('예: export DATA_GO_KR_SERVICE_KEY="발급받은 서비스키"', file=sys.stderr)
        sys.exit(1)

    result = run_query(
        args.corp_name,
        args.category,
        args.spec,
        args.begin_date,
        args.end_date,
        service_key,
        num_of_rows=args.num_of_rows,
    )

    if args.save_raw and result.get("raw_items"):
        raw_output_path = default_output_filename(
            args.corp_name, args.category, args.spec, ext="_원본전체.csv"
        )
        write_raw_output(result["raw_items"], raw_output_path)

    if not result["ok"]:
        sys.exit(2)

    output_path = args.output or default_output_filename(args.corp_name, args.category, args.spec)
    write_output(
        result["camera_items"],
        args.corp_name,
        args.category,
        args.spec,
        args.begin_date,
        args.end_date,
        output_path,
    )


if __name__ == "__main__":
    main()
